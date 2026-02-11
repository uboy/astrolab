from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
import zipfile
from xml.etree import ElementTree

MAX_RESUME_BYTES = 10 * 1024 * 1024
ALLOWED_EXTS = {".docx", ".pdf", ".md"}


def validate_resume_file(filename: str, content_type: str, size_bytes: int) -> None:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTS:
        raise ValueError("Unsupported file type")
    if size_bytes < 0 or size_bytes > MAX_RESUME_BYTES:
        raise ValueError("File too large")


def extract_resume_text(file_bytes: bytes, filename: str, content_type: str) -> str:
    validate_resume_file(filename, content_type, len(file_bytes))
    ext = Path(filename).suffix.lower()
    if ext == ".md":
        return file_bytes.decode("utf-8", errors="ignore")
    if ext == ".docx":
        return _extract_docx_text(file_bytes)
    if ext == ".pdf":
        return _extract_pdf_text(file_bytes)
    raise ValueError("Unsupported file type")


def _extract_docx_text(file_bytes: bytes) -> str:
    # Prefer python-docx for real-world documents with tables/headers.
    try:
        from docx import Document  # type: ignore

        document = Document(BytesIO(file_bytes))
        parts = []
        for p in document.paragraphs:
            if p.text and p.text.strip():
                parts.append(p.text.strip())
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        parts.append(text)
        joined = "\n".join(parts).strip()
        if joined:
            return joined
    except Exception:
        pass

    # Fallback parser for minimal DOCX structures.
    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as zf:
            data = zf.read("word/document.xml")
    except Exception as exc:
        raise ValueError("Invalid docx") from exc

    try:
        root = ElementTree.fromstring(data)
    except Exception as exc:
        raise ValueError("Invalid docx xml") from exc

    texts = []
    for node in root.iter():
        if node.tag.endswith("}t") and node.text:
            texts.append(node.text)
    return "\n".join(texts).strip()


def _extract_pdf_text(file_bytes: bytes) -> str:
    if not file_bytes.startswith(b"%PDF"):
        raise ValueError("Invalid pdf header")

    # Prefer pypdf for normal PDFs with compressed streams.
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(BytesIO(file_bytes))
        chunks = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                chunks.append(page_text.strip())
        joined = "\n".join(chunks).strip()
        if joined:
            return joined
    except Exception:
        pass

    # Fallback parser for very simple PDFs used in tests.
    raw = file_bytes.decode("latin1", errors="ignore")
    streams = re.findall(r"stream\\r?\\n(.*?)\\r?\\nendstream", raw, flags=re.S)
    search_space = "\n".join(streams) if streams else raw

    text_parts = []
    for literal in re.findall(r"\(([^)]*)\)", search_space):
        text_parts.append(_unescape_pdf_text(literal))
    for hex_text in re.findall(r"<([0-9A-Fa-f\\s]+)>", search_space):
        decoded = _decode_pdf_hex_string(hex_text)
        if decoded:
            text_parts.append(decoded)

    text = "\n".join(p for p in text_parts if p).strip()
    if not text:
        raise ValueError("No text found in pdf")
    return text


def _decode_pdf_hex_string(hex_text: str) -> str:
    cleaned = re.sub(r"\\s+", "", hex_text)
    if len(cleaned) % 2 == 1:
        cleaned += "0"
    try:
        raw = bytes.fromhex(cleaned)
    except ValueError:
        return ""
    try:
        return raw.decode("utf-8")
    except Exception:
        return raw.decode("latin1", errors="ignore")


def _unescape_pdf_text(text: str) -> str:
    basic = (
        text.replace(r"\(", "(")
        .replace(r"\)", ")")
        .replace(r"\\", "\\")
        .replace(r"\n", "\n")
        .replace(r"\r", "\r")
        .replace(r"\t", "\t")
    )
    raw = basic.encode("latin1", errors="ignore")
    if b"\\x" in raw:
        raw = _replace_hex_escapes(raw)
    try:
        return raw.decode("utf-8")
    except Exception:
        return raw.decode("latin1", errors="ignore")


def _replace_hex_escapes(data: bytes) -> bytes:
    out = bytearray()
    i = 0
    while i < len(data):
        if i + 3 < len(data) and data[i:i + 2] == b"\\x":
            try:
                out.append(int(data[i + 2:i + 4], 16))
                i += 4
                continue
            except Exception:
                pass
        out.append(data[i])
        i += 1
    return bytes(out)
